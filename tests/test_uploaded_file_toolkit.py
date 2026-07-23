# -*- coding: utf-8 -*-
"""测试 UploadedFileToolkit 的各种文件类型处理。"""

import os
import sys
import tempfile
import zipfile

import pytest

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
os.chdir(ROOT)
sys.path.insert(0, ROOT)

os.environ["PROXY"] = ""

from app.cosight.tool.uploaded_file_toolkit import UploadedFileToolkit


@pytest.fixture
def tk():
    return UploadedFileToolkit()


@pytest.fixture
def tmp(request):
    d = tempfile.mkdtemp()
    yield d
    import shutil
    shutil.rmtree(d, ignore_errors=True)


def w(d, name, content):
    """Write text file."""
    path = os.path.join(d, name)
    with open(path, "w", encoding="utf-8") as f:
        f.write(content)
    return path


def wb(d, name, content):
    """Write binary file."""
    path = os.path.join(d, name)
    with open(path, "wb") as f:
        f.write(content)
    return path


# ========== 文本类 ==========


class TestTextFiles:
    def test_txt(self, tk, tmp):
        result = tk.process_uploaded_file(w(tmp, "a.txt", "Hello World"))
        assert "a.txt" in result
        assert "Hello World" in result

    def test_py(self, tk, tmp):
        result = tk.process_uploaded_file(w(tmp, "a.py", "def hello():\n    pass"))
        assert "a.py" in result
        assert "def hello()" in result

    def test_yaml(self, tk, tmp):
        result = tk.process_uploaded_file(w(tmp, "a.yaml", "key: value"))
        assert "a.yaml" in result
        assert "key:" in result

    def test_md(self, tk, tmp):
        result = tk.process_uploaded_file(w(tmp, "a.md", "# Title\nBody"))
        assert "a.md" in result
        assert "# Title" in result


# ========== 结构化数据 ==========


class TestStructuredData:
    def test_json(self, tk, tmp):
        result = tk.process_uploaded_file(w(tmp, "a.json", '{"x": 1}'))
        assert "a.json" in result
        assert "x" in result

    def test_xml(self, tk, tmp):
        result = tk.process_uploaded_file(w(tmp, "a.xml", "<r><i>v</i></r>"))
        assert "a.xml" in result
        assert "r" in result

    def test_csv(self, tk, tmp):
        result = tk.process_uploaded_file(w(tmp, "a.csv", "a,b\n1,2"))
        assert "a.csv" in result
        assert "a,b" in result


# ========== PDF ==========


class TestPDF:
    def test_pdf_text_extraction(self, tk, tmp):
        pdf_path = os.path.join(tmp, "test.pdf")
        try:
            import fitz
            doc = fitz.open()
            doc.new_page().insert_text((50, 100), "PDF Content Test")
            doc.save(pdf_path)
            doc.close()
        except ImportError:
            pytest.skip("PyMuPDF not installed")

        result = tk.process_uploaded_file(pdf_path)
        assert "test.pdf" in result
        assert "PDF Content Test" in result


# ========== Word ==========


class TestWord:
    def test_docx_structure(self, tk, tmp):
        docx_path = os.path.join(tmp, "test.docx")
        try:
            from docx import Document
            doc = Document()
            doc.add_heading("第一章", 1)
            doc.add_paragraph("正文内容")
            doc.add_heading("1.1 小节", 2)
            doc.add_paragraph("小节内容")
            doc.save(docx_path)
        except ImportError:
            pytest.skip("python-docx not installed")

        result = tk.process_uploaded_file(docx_path)
        assert "test.docx" in result
        assert "第一章" in result
        assert "正文内容" in result
        assert "1.1" in result


# ========== Excel ==========


class TestExcel:
    def test_excel_structure(self, tk, tmp):
        xlsx_path = os.path.join(tmp, "test.xlsx")
        try:
            import pandas as pd
            df = pd.DataFrame({"Name": ["Alice"], "Age": [25]})
            df.to_excel(xlsx_path, index=False, sheet_name="People")
        except ImportError:
            pytest.skip("pandas not installed")

        result = tk.process_uploaded_file(xlsx_path)
        assert "test.xlsx" in result
        assert "People" in result
        assert "Name" in result
        assert "Age" in result
        assert "使用说明" in result


# ========== PPT ==========


class TestPPT:
    def test_pptx_text(self, tk, tmp):
        pptx_path = os.path.join(tmp, "test.pptx")
        try:
            from pptx import Presentation
            prs = Presentation()
            sl = prs.slides.add_slide(prs.slide_layouts[6])
            sl.shapes.title.text_frame.text = "PPT Title"
            prs.save(pptx_path)
        except ImportError:
            pytest.skip("python-pptx not installed")

        result = tk.process_uploaded_file(pptx_path)
        assert "test.pptx" in result
        assert "第 1 页" in result


# ========== 压缩包 ==========


class TestArchive:
    def test_zip_file_tree(self, tk, tmp):
        zip_path = os.path.join(tmp, "test.zip")
        with zipfile.ZipFile(zip_path, "w") as zf:
            zf.writestr("readme.txt", "hello")
            zf.writestr("data/notes.txt", "notes")

        result = tk.process_uploaded_file(zip_path)
        assert "test.zip" in result
        assert "readme.txt" in result
        assert "data/" in result


# ========== 图片 ==========


class TestImage:
    def test_png_metadata(self, tk, tmp):
        img_path = os.path.join(tmp, "test.png")
        try:
            from PIL import Image
            Image.new("RGB", (4, 4), color="red").save(img_path)
        except ImportError:
            pytest.skip("Pillow not installed")

        result = tk.process_uploaded_file(img_path)
        assert "test.png" in result
        assert "ask_question_about_image" in result


# ========== 音频/视频 ==========


class TestMedia:
    def test_mp3_metadata(self, tk, tmp):
        path = wb(tmp, "audio.mp3", b"fake mp3")
        result = tk.process_uploaded_file(path)
        assert "AUDIO" in result
        assert "audio_recognition" in result

    def test_mp4_metadata(self, tk, tmp):
        path = wb(tmp, "video.mp4", b"fake mp4")
        result = tk.process_uploaded_file(path)
        assert "VIDEO" in result


# ========== 边界情况 ==========


class TestEdgeCases:
    def test_file_not_found(self, tk):
        result = tk.process_uploaded_file("/nonexistent/path/file.txt")
        assert "文件不存在" in result

    def test_unsupported_extension(self, tk, tmp):
        # 文件存在但扩展名不支持
        path = w(tmp, "test.xyz", "some content")
        result = tk.process_uploaded_file(path)
        assert "不支持的文件类型" in result
