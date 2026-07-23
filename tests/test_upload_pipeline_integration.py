# -*- coding: utf-8 -*-
"""集成测试：上传 → 复制到工作区 → process_uploaded_file 完整链路。"""

import os
import sys
import json

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
os.chdir(ROOT)
sys.path.insert(0, ROOT)
os.environ["PROXY"] = ""

from cosight_server.deep_research.routers.common import commonRouter
from app.cosight.task.task_manager import TaskManager
from app.cosight.tool.uploaded_file_toolkit import UploadedFileToolkit

_test_app = FastAPI()
_test_app.include_router(commonRouter, prefix="/v1")
client = TestClient(_test_app)


class TestUploadToProcessPipeline:
    """完整链路：上传 → 复制 → 处理。"""

    UPLOAD_DIR = "upload_files"

    def setup_method(self):
        os.makedirs(self.UPLOAD_DIR, exist_ok=True)

    def teardown_method(self):
        if os.path.exists(self.UPLOAD_DIR):
            for f in os.listdir(self.UPLOAD_DIR):
                fp = os.path.join(self.UPLOAD_DIR, f)
                if os.path.isfile(fp):
                    os.remove(fp)

    def _upload_file(self, filename: str, content: bytes, mime: str = "text/plain"):
        """通过 API 上传文件，返回文件 ID。"""
        resp = client.post("/v1/upload", files={"files": (filename, content, mime)})
        assert resp.status_code == 200
        data = resp.json()
        assert data["code"] == 0
        return data["data"]["files"][0]

    # ── 文本文件 ──

    def test_txt_upload_copy_process(self, tmp_path):
        """上传 txt → 复制到工作区 → process_uploaded_file 读取。"""
        f = self._upload_file("hello.txt", b"Hello World!\nSecond line")
        result = TaskManager.copy_uploaded_files_to_workspace(
            upload_ids=[f["id"]], workspace_path=str(tmp_path)
        )
        assert result["success"] is True

        ws_file = os.path.join(str(tmp_path), f["id"])
        assert os.path.exists(ws_file)

        tk = UploadedFileToolkit()
        content = tk.process_uploaded_file(ws_file)
        assert "Hello World" in content
        assert "Second line" in content

    def test_py_upload_copy_process(self, tmp_path):
        """上传 .py 文件。"""
        f = self._upload_file("script.py", b"def add(a,b):\n    return a+b")
        TaskManager.copy_uploaded_files_to_workspace(
            upload_ids=[f["id"]], workspace_path=str(tmp_path)
        )
        tk = UploadedFileToolkit()
        content = tk.process_uploaded_file(os.path.join(str(tmp_path), f["id"]))
        assert "def add" in content

    # ── PDF ──

    def test_pdf_upload_copy_process(self, tmp_path):
        """上传 PDF → 复制 → 提取文字。"""
        try:
            import fitz
        except ImportError:
            pytest.skip("PyMuPDF not installed")

        pdf_path = os.path.join(tmp_path, "temp.pdf")
        doc = fitz.open()
        doc.new_page().insert_text((50, 100), "PDF Integration Test")
        doc.save(pdf_path)
        doc.close()

        with open(pdf_path, "rb") as f:
            pdf_bytes = f.read()
        os.remove(pdf_path)

        f = self._upload_file("test.pdf", pdf_bytes, "application/pdf")
        TaskManager.copy_uploaded_files_to_workspace(
            upload_ids=[f["id"]], workspace_path=str(tmp_path)
        )

        tk = UploadedFileToolkit()
        content = tk.process_uploaded_file(os.path.join(str(tmp_path), f["id"]))
        assert "PDF Integration Test" in content

    # ── Word ──

    def test_docx_upload_copy_process(self, tmp_path):
        """上传 DOCX → 复制 → 提取章节内容。"""
        try:
            from docx import Document
        except ImportError:
            pytest.skip("python-docx not installed")

        docx_path = os.path.join(tmp_path, "temp.docx")
        doc = Document()
        doc.add_heading("测试标题", 1)
        doc.add_paragraph("测试正文内容")
        doc.save(docx_path)

        with open(docx_path, "rb") as f:
            docx_bytes = f.read()
        os.remove(docx_path)

        f = self._upload_file("test.docx", docx_bytes,
                               "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        TaskManager.copy_uploaded_files_to_workspace(
            upload_ids=[f["id"]], workspace_path=str(tmp_path)
        )

        tk = UploadedFileToolkit()
        ws_file = os.path.join(str(tmp_path), f["id"])
        assert os.path.exists(ws_file)
        content = tk.process_uploaded_file(ws_file)
        assert "测试标题" in content
        assert "测试正文内容" in content

    # ── Excel ──

    def test_excel_upload_copy_process(self, tmp_path):
        """上传 Excel → 复制 → 分析结构。"""
        try:
            import pandas as pd
        except ImportError:
            pytest.skip("pandas not installed")

        xlsx_path = os.path.join(tmp_path, "temp.xlsx")
        df = pd.DataFrame({"Name": ["Alice"], "Score": [95]})
        df.to_excel(xlsx_path, index=False, sheet_name="Results")

        with open(xlsx_path, "rb") as f:
            xlsx_bytes = f.read()
        os.remove(xlsx_path)

        f = self._upload_file("data.xlsx", xlsx_bytes,
                               "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        TaskManager.copy_uploaded_files_to_workspace(
            upload_ids=[f["id"]], workspace_path=str(tmp_path)
        )

        tk = UploadedFileToolkit()
        content = tk.process_uploaded_file(os.path.join(str(tmp_path), f["id"]))
        assert "Results" in content
        assert "Name" in content
        assert "Alice" in content

    # ── ZIP ──

    def test_zip_upload_copy_process(self, tmp_path):
        """上传 ZIP → 复制 → 展开文件树。"""
        import zipfile
        import io

        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w") as zf:
            zf.writestr("notes.txt", "zip content")
        buf.seek(0)

        f = self._upload_file("archive.zip", buf.read(), "application/zip")
        TaskManager.copy_uploaded_files_to_workspace(
            upload_ids=[f["id"]], workspace_path=str(tmp_path)
        )

        tk = UploadedFileToolkit()
        content = tk.process_uploaded_file(os.path.join(str(tmp_path), f["id"]))
        assert "notes.txt" in content

    # ── 图片 ──

    def test_png_upload_copy_process(self, tmp_path):
        """上传 PNG → 复制 → 获取元数据。"""
        try:
            from PIL import Image
        except ImportError:
            pytest.skip("Pillow not installed")

        img_path = os.path.join(tmp_path, "temp.png")
        Image.new("RGB", (4, 4), color="red").save(img_path)

        with open(img_path, "rb") as f:
            img_bytes = f.read()
        os.remove(img_path)

        f = self._upload_file("image.png", img_bytes, "image/png")
        TaskManager.copy_uploaded_files_to_workspace(
            upload_ids=[f["id"]], workspace_path=str(tmp_path)
        )

        tk = UploadedFileToolkit()
        content = tk.process_uploaded_file(os.path.join(str(tmp_path), f["id"]))
        # 图片应返回元数据，引导使用 ask_question_about_image
        assert ".png" in content
        assert "ask_question_about_image" in content

    # ── 边界：多个文件 ──

    def test_multiple_files_upload_copy(self, tmp_path):
        """上传多个文件，批量复制到工作区。"""
        f1 = self._upload_file("a.txt", b"file a")
        f2 = self._upload_file("b.txt", b"file b")

        ids = [f1["id"], f2["id"]]
        result = TaskManager.copy_uploaded_files_to_workspace(
            upload_ids=ids, workspace_path=str(tmp_path)
        )
        assert result["copied_count"] == 2

        for fid in ids:
            assert os.path.exists(os.path.join(str(tmp_path), fid))
